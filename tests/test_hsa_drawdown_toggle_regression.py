"""W8b: `hsa_drawdown` became optional, and Sheet 11 had a dangling pointer.

Master plan `docs/superpowers/plans/2026-09-19-optimizers-modules-housing-
master-plan.md` §4, W8b (#330 P6b). Notes:
`docs/superpowers/plans/2026-09-22-w8b-newly-optional-medium-risk-notes.md`.

Two independent switches decide what the HSA drawdown schedule looks like, and
W8b is where they stop being the same thing:

* `hsa_withdrawal_mode == 'optimize'` — a *planning lever*. It decides whether
  an optimizer-proposed schedule exists at all; every other mode leaves the
  `11C. HSA Drawdown` sheet rendering a "not applicable" note.
* the `hsa_drawdown` module toggle — new in W8b. It decides whether that sheet
  is in the workbook at all.

Before W8b the second did not exist, so `build_sheet11` (Roth Conversion) could
render one sentence pointing the reader at `'11C. HSA Drawdown'` gated on the
mode alone and be certain the sheet was there. With the toggle, mode can say
`optimize` on a workbook that does not contain that sheet — a cross-reference
to a missing tab. No golden-master fixture pins this (neither the frozen sample
plan nor `input/demo/` sets `optimize`), so the divergence is constructed here
directly rather than waited for, the same shape W7's engine-gate regression
took for the same reason.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from src.data_io import load_csv, parse_client
from src.module_catalog import CATALOG, OPTIONAL_MODULE_SHEETS, module_enabled
from src.plan_config import ensure_engine_config
from src.planning_engines import project
from src.reporting.sheets_strategy import build_sheet11
from src.reporting.workbook_common import SHEET_LETTER_ORDER, compute_final_sheet_renames

from conftest import TEST_INPUT_DIR
from tests.golden_pricing import FROZEN_GOLDEN_MASTER_PRICES, frozen_holdings_prices

ROOT = Path(__file__).resolve().parents[1]

TOGGLE = "hsa_drawdown"
LEGACY_SHEET = "11C. HSA Drawdown"
POINTER_FRAGMENT = "HSA drawdown schedule that shares this objective"


# ── The registration half (cheap, no projection) ─────────────────────────────

def test_module_is_registered_for_gating_and_rename():
    """The W8a lesson: `optional=True` alone does not stop a sheet building.

    `OPTIONAL_MODULE_SHEETS` is derived from `SHEET_REGISTRY.module_key`, not
    from `CATALOG.optional`, so a module can be "optional" on the switch page
    and still be built unconditionally. Both halves are asserted here.
    """
    assert CATALOG[TOGGLE].optional is True
    assert OPTIONAL_MODULE_SHEETS.get(TOGGLE) == [LEGACY_SHEET], \
        f"{TOGGLE} not registered to {LEGACY_SHEET} in OPTIONAL_MODULE_SHEETS"
    wb = Workbook()
    for ordered in SHEET_LETTER_ORDER.values():
        for stable in ordered:
            wb.create_sheet(stable)
    assert LEGACY_SHEET in compute_final_sheet_renames(wb)


def test_module_enabled_toggle_gates_the_sheet():
    assert module_enabled({"opt": {TOGGLE: True}}, TOGGLE) is True
    assert module_enabled({"opt": {TOGGLE: False}}, TOGGLE) is False


def test_the_soft_dependency_is_declared_on_the_consumer():
    """The pointer lives in Sheet 11, so `roth_conversion_plan` is what
    degrades. W5's call-site sweep enforces that this declaration exists; this
    asserts it names the right thing, which the sweep cannot."""
    losses = dict(CATALOG["roth_conversion_plan"].degrades_without)
    assert TOGGLE in losses, (
        "roth_conversion_plan renders a pointer at the HSA drawdown sheet but "
        "does not declare degrades_without on it")
    assert "HSA drawdown" in losses[TOGGLE]
    # And it must not be a hard prerequisite: Sheet 11's analysis is identical
    # either way, only the one pointer line differs.
    assert TOGGLE not in CATALOG["roth_conversion_plan"].requires_outputs


# ── The pointer half (one projection, shared across the three cases) ─────────

@pytest.fixture(scope="module")
def _projected():
    c = ensure_engine_config(
        parse_client(load_csv(TEST_INPUT_DIR / "client_data.csv"), ""), source="test")
    c['mc_sims'] = 8
    c['mc_sensitivity_sims'] = 1
    with frozen_holdings_prices(FROZEN_GOLDEN_MASTER_PRICES):
        rows = project(c)
    return c, rows


def _sheet11_text(c, rows, *, mode, toggle):
    trial = dict(c)
    trial['hsa_withdrawal_mode'] = mode
    trial['opt'] = dict(trial.get('opt') or {})
    trial['opt'][TOGGLE] = toggle
    ws = Workbook().create_sheet("t")
    build_sheet11(ws, trial, rows)
    return " ".join(
        str(cell.value) for row in ws.iter_rows()
        for cell in row if cell.value is not None
    )


def test_optimize_mode_with_the_module_on_still_points_at_the_sheet(_projected):
    """The control. Without this the two assertions below could both pass on a
    sheet that simply stopped rendering the pointer at all."""
    c, rows = _projected
    assert POINTER_FRAGMENT in _sheet11_text(c, rows, mode='optimize', toggle=True)


def test_optimize_mode_with_the_module_off_drops_the_pointer(_projected):
    """The defect W8b would otherwise have introduced: mode says `optimize`, so
    the pre-W8b gate renders the pointer, but the toggle means
    `11C. HSA Drawdown` is not in the workbook for it to point at."""
    c, rows = _projected
    assert POINTER_FRAGMENT not in _sheet11_text(c, rows, mode='optimize', toggle=False)


def test_a_non_optimize_mode_never_renders_the_pointer_either_way(_projected):
    """The mode gate is unchanged by W8b -- the toggle was ANDed onto it, not
    substituted for it. Both fixtures ship `smooth_window`, so this is also the
    case every golden master actually exercises."""
    c, rows = _projected
    for toggle in (True, False):
        assert POINTER_FRAGMENT not in _sheet11_text(
            c, rows, mode='smooth_window', toggle=toggle)

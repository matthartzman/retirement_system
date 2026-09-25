"""#332 W-F Task F3 (design 2026-09-24 §1.3, §9.1): the workbook's Plan Data
section-index table (built from _SECTION_META / WORKBOOK_SECTION_LAYOUT)
gains a Topic column so the answer-type-organized workbook also shows each
sheet's catalog domain -- the same facet the left nav and Plan Features use.
"""
import pytest
from openpyxl import load_workbook

from src.module_catalog import CATALOG, SHEET_REGISTRY, WHOLE_PLAN
from src.reporting.workbook_common import sheet_topic


def test_sheet_topic_matches_the_owning_modules_domain():
    for name, spec in SHEET_REGISTRY.items():
        if spec.module_key:
            assert sheet_topic(name) == CATALOG[spec.module_key].domain, name
        else:
            assert sheet_topic(name) == WHOLE_PLAN, name


@pytest.mark.slow
def test_plan_data_section_index_has_a_topic_column(built_workbook_path):
    wb = load_workbook(built_workbook_path, read_only=False, data_only=False)
    ws = next(s for s in wb.worksheets if s.title.split(".", 1)[-1].strip() == "Plan Data" or s.title == "Plan Data")
    header_row = next(
        r for r in range(1, ws.max_row + 1)
        if ws.cell(row=r, column=1).value == "Section"
    )
    headers = [ws.cell(row=header_row, column=c).value for c in range(1, 5)]
    assert headers == ["Section", "Sheet", "Topic", "Purpose"]

    known_topics = set(sheet_topic(name) for name in SHEET_REGISTRY)
    seen_any = False
    for r in range(header_row + 1, ws.max_row + 1):
        section = ws.cell(row=r, column=1).value
        if not section:
            break
        topic = ws.cell(row=r, column=3).value
        assert topic in known_topics, (r, topic)
        seen_any = True
    assert seen_any

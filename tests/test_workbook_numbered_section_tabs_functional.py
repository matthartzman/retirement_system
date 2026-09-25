import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.slow


def test_workbook_uses_numbered_sections_and_lettered_children(built_workbook_path):
    assert built_workbook_path.exists(), 'Generated workbook is missing'
    wb = load_workbook(built_workbook_path, read_only=False, data_only=False)
    visible = [ws.title for ws in wb.worksheets if ws.sheet_state == 'visible']
    # #332 W-F Task F6 (design 2026-09-24 §9.1): within each numbered section,
    # sheets now order by Topic (module_catalog.DOMAINS order) first and the
    # legacy letter_rank second, so the tab strip shows both facets: answer
    # type (the section number) and Topic (the letter sequence). Letters are
    # still derived, densely, from this order -- a module-gated sheet absent
    # from this fixture's plan simply compresses the remaining letters. A
    # sheet with no owning catalog module (Executive Summary, Net Worth,
    # Asset Allocation, Plan Data, ...) sorts last within its section, as
    # workbook_common.sheet_topic() returns "Whole Plan" for it.
    expected = [
        '1. Reports',
        '1A. Spending Summary',
        '1B. Lifetime Taxes',
        '1C. Executive Summary',
        '1D. Net Worth',
        '1E. Cash Flow',
        '1F. Balance Sheet',
        '1G. Charts',
        '1H. Current vs. Proposed',
        '1I. Planning Levers',
        '2. Optimizers',
        '2A. Social Security',
        '2B. Housing Comparison',
        '2C. Withdrawal Sequencing',
        '2D. Asset Location',
        '2E. Roth Conversion',
        # HSA Drawdown always sits right after Roth Conversion (shares its
        # objective, per sheets_strategy.py) -- always-on, not module-gated,
        # so it is never absent.
        '2F. HSA Drawdown',
        '2G. Charitable Giving',
        '2H. Tax-Loss Harvesting',
        '2I. Gain Harvesting',
        '2J. Estate & Legacy Planning',
        # Asset Allocation has no owning catalog module (it merges
        # allocation-policy output across several inputs), so it is
        # Topic "Whole Plan" and sorts last in this section regardless of
        # letter_rank.
        '2K. Asset Allocation',
        '3. Comparisons',
        '3A. State Residency',
        '3B. S-Corp vs LLC',
        '3C. Scenario Analysis',
        '4. Risks',
        '4A. Monte Carlo',
        '4B. Survivor',
        # W3 (#329 O10): LTC Stress Test is split back out of the merged Life
        # Insurance sheet into its own tab under 4.1 stress tests.
        '4C. LTC Stress Test',
        # #329 §1.2/§3.3 (W9): Divorce/QDRO would sit here (Insurance & Care,
        # between Survivor and LTC Stress Test) if enabled -- this fixture
        # does not force-enable it (off by default), matching the "newer
        # default-off modules are intentionally NOT force-enabled" comment
        # on RETIREMENT_SYSTEM_FORCE_ENABLE_MODULES above, so letters
        # compress and Life Insurance Need stays 4D. Only Life Insurance Need
        # is on in this fixture's plan (existing_life_insurance/
        # disability_income_insurance/property_casualty_umbrella are off).
        '4D. Life Insurance Need',
        '5. Reference',
        '5A. RMD Audit',
        '5B. Tax Capacity',
        # Plan Data, Assumptions, Account Reconciliation, Quality Control,
        # Methodology and Glossary all have no owning catalog module -- Topic
        # "Whole Plan" -- and sort last, in their original letter_rank order.
        '5C. Plan Data',
        '5D. Assumptions',
        '5E. Account Reconciliation',
        '5F. Quality Control',
        '5G. Methodology',
        '5H. Glossary',
    ]
    assert visible[: len(expected)] == expected
    assert '5C. Plan Data' in visible
    assert 'Reports' not in visible
    assert 'Risk' not in visible
    assert 'Optimizers' not in visible
    assert 'System Configuration' not in visible


def test_summary_tabs_reference_child_tabs(built_workbook_path):
    wb = load_workbook(built_workbook_path, read_only=False, data_only=False)
    summary_expected = {
        '1. Reports': ['1C. Executive Summary', '1D. Net Worth', '1E. Cash Flow', '1F. Balance Sheet', '1G. Charts'],
        '2. Optimizers': ['2E. Roth Conversion', '2F. HSA Drawdown', '2K. Asset Allocation', '2C. Withdrawal Sequencing', '2A. Social Security', '2J. Estate & Legacy Planning'],
        '3. Comparisons': ['3A. State Residency', '3B. S-Corp vs LLC', '3C. Scenario Analysis'],
        '4. Risks': ['4A. Monte Carlo', '4B. Survivor', '4C. LTC Stress Test', '4D. Life Insurance Need'],
        '5. Reference': ['5C. Plan Data', '5D. Assumptions', '5E. Account Reconciliation', '5F. Quality Control', '5A. RMD Audit', '5G. Methodology', '5H. Glossary'],
    }
    for sheet, children in summary_expected.items():
        ws = wb[sheet]
        values = [cell.value for row in ws.iter_rows(min_row=1, max_row=20, min_col=1, max_col=8) for cell in row]
        for child in children:
            assert child in values


def test_strategy_scorp_ltc_and_asset_location_merges_are_present(built_workbook_path):
    wb = load_workbook(built_workbook_path, read_only=False, data_only=False)
    exec_text = ' '.join(str(c.value or '') for row in wb['1C. Executive Summary'].iter_rows() for c in row)
    scorp_text = ' '.join(str(c.value or '') for row in wb['3B. S-Corp vs LLC'].iter_rows() for c in row)
    allocation_text = ' '.join(str(c.value or '') for row in wb['2K. Asset Allocation'].iter_rows() for c in row)
    ltc_text = ' '.join(str(c.value or '') for row in wb['4C. LTC Stress Test'].iter_rows() for c in row)
    assert 'WITHDRAWAL SEQUENCE STRATEGY' in exec_text
    assert 'S-CORP vs LLC' in scorp_text and 'LLC / Sole-Prop' in scorp_text
    assert 'ASSET-LOCATION OPTIMIZER' in allocation_text
    assert 'LONG-TERM-CARE STRESS TEST' in ltc_text

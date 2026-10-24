"""Shared helpers for the Rental Properties workbook tests (WP12.7).

``case`` projects a ``tests/rental_fixtures.py`` case once per process; ``built_workbooks`` builds
the real workbook with the feature off and on once per session (slow tests only).
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook

import src.module_catalog as mc
from src.reporting.sheets_rental import build_sheet_rental
from tests.rental_fixtures import CASES, base_config, project, rental_config

ROOT = Path(__file__).resolve().parents[1]
ALL_CASES = list(CASES)

_CACHE: dict[str, tuple[dict, list]] = {}
_BUILT: dict[str, Path] = {}


def case(name: str):
    """(config, rows) of a fixture case with the feature on; ``off`` is the sample plan alone."""
    if name not in _CACHE:
        c = base_config() if name == "off" else rental_config(name)
        _CACHE[name] = (c, project(c))
    return _CACHE[name]


def texts(ws) -> list[str]:
    return [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None]


def header_col(ws, header_row: int, text: str) -> int:
    for col in range(1, ws.max_column + 1):
        if ws.cell(row=header_row, column=col).value == text:
            return col
    raise AssertionError(f"header {text!r} not found in row {header_row}")


def row_labeled(ws, label: str, col: int = 1) -> int:
    for row in ws.iter_rows():
        v = row[col - 1].value
        if isinstance(v, str) and v.strip() == label:
            return row[0].row
    raise AssertionError(f"row {label!r} not found")


def rental_sheet(name: str):
    c, rows = case(name)
    ws = Workbook().active
    rep = build_sheet_rental(ws, c, rows)
    return ws, rep, c, rows


def _build(tmp: Path, rows: list | None) -> Path:
    """Build the workbook from a copy of the session workspace; with ``rows`` the rentals are
    stored and the feature switched on first."""
    ws_root = Path(os.environ["RETIREMENT_SYSTEM_WORKSPACE_ROOT"])
    work = tmp / "workspace"
    shutil.copytree(ws_root, work)
    if rows is not None:
        from src.stores import PlanStore
        with PlanStore.open(work / "plan.rpx") as store:
            store.rental.replace_all(rows)
            mc.set_feature(store, "rental_property_income", True)
    env = os.environ.copy()
    env.update(RETIREMENT_SYSTEM_WORKSPACE_ROOT=str(work), RETIREMENT_SYSTEM_PLAN_DB=str(work / "plan.rpx"), RETIREMENT_SYSTEM_OUTPUT_DIR=str(tmp / "out"),
               RETIREMENT_SYSTEM_APP_MODE="LOCAL", RETIREMENT_SYSTEM_WORKSPACE_ID="local",
               RETIREMENT_SYSTEM_DISABLE_LIVE_PRICE_PROVIDERS="1", RETIREMENT_MC_SIMS="16", RETIREMENT_MC_SENSITIVITY_SIMS="3")
    result = subprocess.run([sys.executable, "tools/build_workbook.py"], cwd=ROOT, env=env, text=True, capture_output=True, timeout=600)
    assert result.returncode == 0, result.stdout[-2000:] + result.stderr[-2000:]
    found = list((tmp / "out").rglob("retirement_plan.xlsx"))
    assert found, result.stdout[-2000:]
    return found[0]


def built_workbooks(tmp_path_factory) -> tuple[Path, Path]:
    """(feature off, feature on with the loss-making fixture) workbook paths, built once per session."""
    if not _BUILT:
        _BUILT["off"] = _build(tmp_path_factory.mktemp("rental_off"), None)
        _BUILT["on"] = _build(tmp_path_factory.mktemp("rental_on"), CASES["standalone_loss"].rows)
    return _BUILT["off"], _BUILT["on"]

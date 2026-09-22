"""W10a: the Roth optimizer result reaches the UI through plan_summary.json.

Master plan `docs/superpowers/plans/2026-09-19-optimizers-modules-housing-
master-plan.md` §4, W10a (#329 P5). Notes:
`docs/superpowers/plans/2026-09-22-w10a-roth-result-panel-notes.md`.

#329 §4.5 path 1 -- "read from the last build" -- says the Roth panel should
display `roth_optimization` / `roth_strategy_result`, which are already
attached to the plan result, rather than re-running anything. The one thing
that was missing is a path from that contract to the browser:
`plan_summary.json` (served by `/api/summary`, and the payload `/api/build`
returns as `kpi`) is the artifact the UI already reads for "what did the last
build conclude", so the result rides along there.

Two independent claims:

* the payload projection itself -- shape, the candidate cap, and what is
  deliberately left out -- which is fast and needs no build; and
* a real subprocess build proving the candidate table the panel will draw and
  the candidate table Sheet 11 prints carry the *same* ranks, labels and
  relative scores. That second claim is the reason the 0-100 normalization
  moved into `summary_figures` in the first place, and asserting it against
  the real workbook is the only way to know the two surfaces did not drift.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

from src.reporting.summary_figures import (
    ROTH_CANDIDATE_DISPLAY_LIMIT,
    roth_candidate_objective_values,
    roth_candidate_relative_scores,
    roth_strategy_result_payload,
)

ROOT = Path(__file__).resolve().parents[1]
SHEET_SUFFIX = "Roth Conversion"


def _config(candidate_count=3, **extra):
    candidates = [
        {
            "label": f"Candidate {i}",
            "policy": "optimize_terminal_tax",
            "score": 100.0 - 10.0 * i,
            "total_objective_score": 100.0 - 10.0 * i,
            "lifetime_tax": 1000.0 + i,
            "after_tax_terminal_nw": 50_000.0 - i,
            "total_conversion": 250.0 * i,
        }
        for i in range(candidate_count)
    ]
    c = {
        "roth_optimization": {
            "selected_label": "Candidate 0",
            "selected_policy": "optimize_terminal_tax",
            "objective_mode": "BALANCED_RETIREMENT",
            "target_bracket": 0.24,
            "auto_optimized": True,
            "candidates": candidates,
        }
    }
    c.update(extra)
    return c


# ---------------------------------------------------------------------------
# The payload projection
# ---------------------------------------------------------------------------


def test_no_optimizer_result_returns_none_rather_than_an_empty_shell():
    """`summary_figures`' own contract: return None when the underlying
    analysis is unavailable, so the caller omits the panel instead of drawing
    an empty table."""
    assert roth_strategy_result_payload({}) is None
    assert roth_strategy_result_payload({"roth_optimization": {}}) is None


def test_candidates_are_capped_and_the_full_count_is_still_reported():
    payload = roth_strategy_result_payload(_config(candidate_count=25))
    assert len(payload["candidates"]) == ROTH_CANDIDATE_DISPLAY_LIMIT
    # Without this the panel would say "10 candidates" when 25 were scored.
    assert payload["candidate_count"] == 25


def test_year_by_year_binding_constraints_are_left_out():
    """Workbook-depth diagnostics, one row per projected year. The panel shows
    the candidate comparison; shipping the trace would inflate every
    plan_summary.json for a table nothing on screen draws."""
    c = _config()
    c["roth_strategy_result"] = {
        "selected_strategy_name": "Candidate 0",
        "binding_constraints_by_year": [{"year": 2030, "conversion": 1.0}],
        "candidates": c["roth_optimization"]["candidates"],
    }
    payload = roth_strategy_result_payload(c)
    assert "binding_constraints_by_year" not in payload


def test_the_contract_wins_over_the_raw_optimizer_dict():
    """Same precedence Sheet 11 and `roth_strategy_candidates` already use."""
    c = _config()
    c["roth_strategy_result"] = {
        "selected_strategy_name": "From the contract",
        "why_selected": "Because the contract said so.",
        "candidates": c["roth_optimization"]["candidates"],
    }
    payload = roth_strategy_result_payload(c)
    assert payload["selected_strategy_name"] == "From the contract"
    assert payload["why_selected"] == "Because the contract said so."


def test_relative_scores_span_zero_to_one_hundred():
    scores = roth_candidate_relative_scores([10.0, 5.0, 0.0])
    assert scores[0] == pytest.approx(100.0)
    assert scores[-1] == pytest.approx(0.0)


def test_a_set_with_no_spread_scores_one_hundred_rather_than_dividing_by_zero():
    assert roth_candidate_relative_scores([7.0, 7.0]) == [100.0, 100.0]
    assert roth_candidate_relative_scores([7.0]) == [100.0]
    assert roth_candidate_relative_scores([]) == []


def test_objective_values_read_the_newest_shape_first():
    """`total_objective_score` is the contract's name; `score` is the raw
    optimizer row's. Both reach this function -- see `roth_strategy_candidates`."""
    assert roth_candidate_objective_values([{"total_objective_score": 3.0, "score": 9.0}]) == [3.0]
    assert roth_candidate_objective_values([{"score": 9.0}]) == [9.0]


# ---------------------------------------------------------------------------
# The real build
# ---------------------------------------------------------------------------


def _build(tmp_path_factory):
    out_dir = tmp_path_factory.mktemp("roth_result_panel")
    env = os.environ.copy()
    env["RETIREMENT_SYSTEM_OUTPUT_DIR"] = str(out_dir)
    env["RETIREMENT_SYSTEM_APP_MODE"] = "LOCAL"
    env["RETIREMENT_SYSTEM_WORKSPACE_ID"] = "local"
    env["RETIREMENT_SYSTEM_DISABLE_LIVE_PRICE_PROVIDERS"] = "1"
    env["RETIREMENT_MC_SIMS"] = "16"
    env["RETIREMENT_MC_SENSITIVITY_SIMS"] = "3"
    result = subprocess.run(
        [sys.executable, "tools/build_workbook.py"],
        cwd=ROOT, env=env, text=True, capture_output=True, timeout=900,
    )
    return out_dir, result


def _sheet11_candidate_rows(book_path):
    """The Sheet 11 candidate table, read back as (rank, label, score) rows."""
    from openpyxl import load_workbook

    wb = load_workbook(book_path, read_only=True, data_only=True)
    try:
        name = next(n for n in wb.sheetnames if n.endswith(SHEET_SUFFIX))
        ws = wb[name]
        grid = [[cell.value for cell in row] for row in ws.iter_rows()]
    finally:
        wb.close()
    header_at = None
    for idx, row in enumerate(grid):
        if row and row[0] == "Rank" and "Candidate" in [str(v) for v in row[:3]]:
            header_at = idx
            break
    assert header_at is not None, "Sheet 11's candidate table header was not found"
    rows = []
    for row in grid[header_at + 1:]:
        if not row or not isinstance(row[0], (int, float)):
            break
        rows.append((int(row[0]), str(row[1]), float(row[3])))
    return rows


@pytest.mark.slow
def test_a_real_build_writes_the_panels_payload_into_plan_summary(tmp_path_factory):
    out_dir, result = _build(tmp_path_factory)
    tail = (result.stdout + result.stderr)[-4000:]
    assert result.returncode == 0, f"build failed:\n{tail}"

    summary = json.loads((out_dir / "plan_summary.json").read_text(encoding="utf-8"))
    payload = summary.get("roth_strategy_result")
    assert payload, "plan_summary.json carries no roth_strategy_result for the UI to read"
    assert payload["selected_strategy_name"], "a selected strategy is the panel's headline"
    assert payload["candidates"], "the candidate comparison is the panel's whole point"
    assert payload["candidate_count"] >= len(payload["candidates"])
    # #329 §4.6's "applied" comparison (W10c) and the panel's own copy both
    # need to name the policy the plan is actually running.
    assert payload["selected_policy"]

    sheet_rows = _sheet11_candidate_rows(out_dir / "retirement_plan.xlsx")
    assert sheet_rows, "Sheet 11 printed no candidates to compare against"
    panel_rows = [
        (c["rank"], c["label"], c["relative_score"]) for c in payload["candidates"]
    ]
    assert len(panel_rows) == len(sheet_rows)
    for panel, sheet in zip(panel_rows, sheet_rows):
        assert panel[0] == sheet[0]
        assert panel[1] == sheet[1]
        # The whole reason the normalization moved into summary_figures: one
        # candidate, one score, on both surfaces.
        assert panel[2] == pytest.approx(sheet[2], abs=1e-6)

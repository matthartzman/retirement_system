"""W10c: the Social Security claim-age sweep's result reaches the apply path.

Master plan `docs/superpowers/plans/2026-09-19-optimizers-modules-housing-
master-plan.md` §4, W10c (#329 P6). Notes:
`docs/superpowers/plans/2026-09-22-w10c-apply-to-plan-notes.md`.

Social Security is #329 §4.3's *scalar adoption* case: the whole
recommendation is one claim age per person, so applying it is a two-row patch
on rows that already exist. The one thing missing was a path from the sweep to
the browser. `build_sheet10()` already computes it and hands it to Sheet 1;
this projects that same dict onto `plan_summary.json` beside W10a's
`roth_strategy_result`, which is the artifact `/api/summary` already serves.

Nothing here re-runs a sweep. Two claims:

* the payload projection -- shape, what it reports when the sweep is absent or
  partial, and the Member 1 / Member 2 mapping the patch rows are keyed by,
  which is fast and needs no build; and
* a real subprocess build proving the key is actually written, and that the
  ages it carries are the same ones Sheet 10 recommends. The second claim is
  the one that catches the projection silently drifting from the sweep.
"""
from __future__ import annotations

import json
import os
import re
import subprocess
import sys
from pathlib import Path

import pytest

from src.reporting.summary_figures import social_security_timing_payload

ROOT = Path(__file__).resolve().parents[1]
SHEET_SUFFIX = "Social Security"


def _sweep(**over):
    """A build_sheet10() return value, trimmed to the keys this reads."""
    best = {
        "h_age": 70, "w_age": 67, "rank_score": 100, "objective_value": 4_100_000.0,
        "after_tax_terminal_nw": 3_900_000.0, "survivor_period_ss_income": 200_000.0,
        "lcv": 5_000_000.0, "lifetime_ss": 1_250_000.0,
    }
    current = {
        "h_age": 67, "w_age": 67, "rank_score": 82, "objective_value": 3_950_000.0,
        "after_tax_terminal_nw": 3_800_000.0, "lcv": 4_900_000.0,
    }
    sweep = {
        "best": best, "current": current,
        "scenarios": [best, current] + [dict(best) for _ in range(31)],
        "h_current": 67, "w_current": 67,
        "h_label": "Pat", "w_label": "Sam",
        "all_infeasible": False,
    }
    sweep.update(over)
    return sweep


# ---------------------------------------------------------------------------
# The projection
# ---------------------------------------------------------------------------


def test_the_recommended_pair_is_reported_per_member():
    payload = social_security_timing_payload(_sweep())
    assert payload["recommended_member_1_claim_age"] == 70
    assert payload["recommended_member_2_claim_age"] == 67


def test_h_is_member_one_and_w_is_member_two():
    """The mapping the patch rows are keyed by, asserted rather than assumed.

    `src/data_io.py` reads `h_ss_pia` from the `Member 1` subsection and
    `w_ss_pia` from `Member 2`; the frontend patch writes the recommended
    `h_age` into `Member 1`'s `claim_date` row on the strength of that. A
    silent swap here would apply each person's answer to the other, which no
    shape assertion would catch -- so this reads the mapping back out of
    data_io's own source rather than restating it.
    """
    payload = social_security_timing_payload(_sweep(best={"h_age": 62, "w_age": 70}))
    assert payload["recommended_member_1_claim_age"] == 62
    assert payload["recommended_member_2_claim_age"] == 70

    src = (ROOT / "src" / "data_io.py").read_text(encoding="utf-8")
    h_pia = re.search(r"c\['h_ss_pia'\]\s*=.*?'Social Security','([^']+)'", src)
    w_pia = re.search(r"c\['w_ss_pia'\]\s*=.*?'Social Security','([^']+)'", src)
    assert h_pia and w_pia, "data_io.py's h_/w_ Social Security mapping was not found"
    assert h_pia.group(1) == "Member 1"
    assert w_pia.group(1) == "Member 2"


def test_the_configured_pair_comes_from_the_sweeps_own_h_current():
    """`current` is None whenever the configured pair was never scored.

    The coarse-then-refine pass only scores part of the grid, so the
    configured pair can legitimately be missing from `scenarios`. The
    configured AGES must still be reported -- they are what the panel compares
    the recommendation against -- so they come from h_current/w_current, which
    are always set, not from the `current` row, which is not.
    """
    payload = social_security_timing_payload(_sweep(current=None))
    assert payload["configured_member_1_claim_age"] == 67
    assert payload["configured_member_2_claim_age"] == 67
    assert payload["configured_was_scored"] is False
    assert payload["configured_objective_value"] is None


def test_the_configured_pairs_figures_are_reported_when_it_was_scored():
    payload = social_security_timing_payload(_sweep())
    assert payload["configured_was_scored"] is True
    assert payload["configured_objective_value"] == 3_950_000.0
    assert payload["configured_lcv"] == 4_900_000.0


def test_a_recommendation_the_plan_already_runs_says_so():
    payload = social_security_timing_payload(_sweep(h_current=70, w_current=67))
    assert payload["recommendation_matches_plan"] is True


def test_a_recommendation_that_differs_says_so():
    assert social_security_timing_payload(_sweep())["recommendation_matches_plan"] is False


def test_no_sweep_returns_none_rather_than_an_empty_shell():
    """The Social Security optimizer module being off is the normal way this
    happens (`ss_sweep` stays None in workbook_builder), and an absent key is
    how the UI is told there is nothing to apply."""
    assert social_security_timing_payload(None) is None
    assert social_security_timing_payload({}) is None
    assert social_security_timing_payload({"best": {}}) is None


def test_a_single_person_household_reports_no_member_two_age():
    payload = social_security_timing_payload(_sweep(best={"h_age": 70, "w_age": None}))
    assert payload["recommended_member_1_claim_age"] == 70
    assert payload["recommended_member_2_claim_age"] is None


def test_the_full_scenario_grid_is_left_out_but_its_size_is_reported():
    """Apply-to-plan needs the winner, not 81 pairs of workbook diagnostics."""
    payload = social_security_timing_payload(_sweep())
    assert "scenarios" not in payload
    assert payload["pairs_scored"] == 33


def test_an_all_infeasible_sweep_is_flagged_rather_than_presented_as_clean():
    """Every pair failing the essential-funding gate means the recommendation
    is a least-bad fallback. Applying it is legitimate; saying nothing about
    it is not."""
    assert social_security_timing_payload(_sweep())["all_pairs_infeasible"] is False
    assert social_security_timing_payload(_sweep(all_infeasible=True))["all_pairs_infeasible"] is True


def test_the_payload_is_json_serializable():
    """It rides on plan_summary.json, so a stray numpy scalar would break the
    whole artifact, not just this key."""
    json.dumps(social_security_timing_payload(_sweep()))


def test_no_applied_flag_is_ever_written_server_side():
    """#329 §4.6: applied state is computed client-side by comparing the
    result against the live rows. A server-side flag would go stale the
    moment anyone edited a row by hand, which is the failure the design names
    explicitly."""
    payload = social_security_timing_payload(_sweep())
    assert not [k for k in payload if k == "applied" or k.endswith("_applied")]


def test_the_builder_returns_the_configured_pair_for_this_projection():
    """`h_current`/`w_current` exist on build_sheet10's return value because
    this payload needs them; a test pins that so the keys are not dropped as
    unused."""
    src = (ROOT / "src" / "reporting" / "sheets_strategy.py").read_text(encoding="utf-8")
    assert "'h_current': h_current, 'w_current': w_current," in src


def test_workbook_builder_wires_the_payload_into_plan_summary():
    src = (ROOT / "src" / "reporting" / "workbook_builder.py").read_text(encoding="utf-8")
    assert "'social_security_timing_result': None," in src
    assert "social_security_timing_payload(ss_sweep, c)" in src


# ---------------------------------------------------------------------------
# The real build
# ---------------------------------------------------------------------------


@pytest.mark.slow
def test_a_real_build_writes_the_apply_payload_into_plan_summary(tmp_path_factory):
    out_dir = tmp_path_factory.mktemp("ss_apply_payload")
    env = os.environ.copy()
    env["RETIREMENT_SYSTEM_OUTPUT_DIR"] = str(out_dir)
    env["RETIREMENT_SYSTEM_APP_MODE"] = "LOCAL"
    env["RETIREMENT_SYSTEM_WORKSPACE_ID"] = "local"
    env["RETIREMENT_SYSTEM_DISABLE_LIVE_PRICE_PROVIDERS"] = "1"
    env["RETIREMENT_MC_SIMS"] = "16"
    env["RETIREMENT_MC_SENSITIVITY_SIMS"] = "3"
    result = subprocess.run(
        [sys.executable, "tools/build_workbook.py"],
        cwd=ROOT, env=env, text=True, capture_output=True, timeout=1800,
    )
    tail = (result.stdout + result.stderr)[-4000:]
    assert result.returncode == 0, f"build failed:\n{tail}"

    summary = json.loads((out_dir / "plan_summary.json").read_text(encoding="utf-8"))
    payload = summary.get("social_security_timing_result")
    assert payload, "plan_summary.json carries no social_security_timing_result"
    assert 62 <= payload["recommended_member_1_claim_age"] <= 70
    assert payload["pairs_scored"] > 0

    # The payload and Sheet 1's own "Recommended ... Claim Age" headline are
    # both projections of the same `best`; if they ever disagree the panel is
    # telling the user to apply something the workbook does not recommend.
    from openpyxl import load_workbook

    book = next(out_dir.glob("*.xlsx"))
    wb = load_workbook(book, read_only=True, data_only=True)
    try:
        ws = wb["1. Executive Summary"]
        headline = {}
        for row in ws.iter_rows():
            vals = [c.value for c in row]
            for idx, v in enumerate(vals[:-1]):
                if isinstance(v, str) and v.startswith("Recommended ") and v.endswith("Claim Age"):
                    headline[v] = vals[idx + 1]
    finally:
        wb.close()
    assert headline, "Sheet 1 printed no Recommended Claim Age headline to compare against"
    ages = {int(v) for v in headline.values() if isinstance(v, (int, float))}
    assert payload["recommended_member_1_claim_age"] in ages, (
        f"payload says {payload['recommended_member_1_claim_age']}, Sheet 1 says {headline}"
    )

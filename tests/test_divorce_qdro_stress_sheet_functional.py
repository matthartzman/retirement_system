"""W9: Divorce/QDRO gains a workbook sheet.

Master plan `docs/superpowers/plans/2026-09-19-optimizers-modules-housing-
master-plan.md` §4, W9 (#329 P4 + #330). Notes:
`docs/superpowers/plans/2026-09-22-w9-ui-section-registry-notes.md`.

`divorce_qdro` was `sheet=None` -- a UI-only stress test with no workbook
counterpart, "the mirror image of the workbook-only optimizers" (#329 §1.2).
`build_sheet39` gives it one, reusing the exact re-projection Sheet 16
(Scenario Analysis)'s own "Divorce/QDRO Asset Split" comparison row already
computes (`divorce_split_yr`/`divorce_split_pct` overrides on `run_scenario`).

Two independent claims: the registration half (fast, no build), and a real
subprocess build proving the toggle gates the sheet rather than merely
surviving with it always present or always absent -- the same shape W8b's
housing/HSA regressions took for the same reason.
"""
from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

import pytest

from src.module_catalog import CATALOG, OPTIONAL_MODULE_SHEETS

ROOT = Path(__file__).resolve().parents[1]
TOGGLE = "divorce_qdro"
SHEET_NAME = "39. Divorce QDRO Stress Test"


def test_module_is_registered_for_gating():
    """`optional=True` alone would not stop the sheet building (W8a's
    finding) -- `OPTIONAL_MODULE_SHEETS` (derived from `SHEET_REGISTRY.
    module_key`) is what the build gate actually reads."""
    assert CATALOG[TOGGLE].optional is True
    assert OPTIONAL_MODULE_SHEETS.get(TOGGLE) == [SHEET_NAME], \
        f"{TOGGLE} not registered to {SHEET_NAME} in OPTIONAL_MODULE_SHEETS"


def test_catalog_entry_has_its_sheet_and_tab():
    m = CATALOG[TOGGLE]
    assert m.sheet == SHEET_NAME
    assert m.tab
    assert m.kind == "stress_test"


def _build_workbook(tmp_path_factory, *, enabled: bool):
    out_dir = tmp_path_factory.mktemp(f"divorce_qdro_{'on' if enabled else 'off'}")
    env = os.environ.copy()
    env["RETIREMENT_SYSTEM_OUTPUT_DIR"] = str(out_dir)
    env["RETIREMENT_SYSTEM_APP_MODE"] = "LOCAL"
    env["RETIREMENT_SYSTEM_WORKSPACE_ID"] = "local"
    env["RETIREMENT_SYSTEM_DISABLE_LIVE_PRICE_PROVIDERS"] = "1"
    env["RETIREMENT_MC_SIMS"] = "16"
    env["RETIREMENT_MC_SENSITIVITY_SIMS"] = "3"
    for key in ("RETIREMENT_SYSTEM_FORCE_ALL_MODULES",
                "RETIREMENT_SYSTEM_FORCE_ENABLE_MODULES",
                "RETIREMENT_SYSTEM_FORCE_DISABLE_MODULES"):
        env.pop(key, None)
    env["RETIREMENT_SYSTEM_FORCE_ENABLE_MODULES" if enabled
        else "RETIREMENT_SYSTEM_FORCE_DISABLE_MODULES"] = TOGGLE
    result = subprocess.run(
        [sys.executable, "tools/build_workbook.py"],
        cwd=ROOT, env=env, text=True, capture_output=True, timeout=600,
    )
    return out_dir, result


@pytest.mark.slow
@pytest.mark.parametrize("enabled", [True, False])
def test_full_workbook_build_succeeds_with_the_module_on_and_off(tmp_path_factory, enabled):
    """The off case asserts the sheet is gone; the on case asserts it is
    present with real content -- not just that the build survived either way.

    Matched on the stable display suffix rather than a letter (`4D.` today):
    that letter shifts whenever another Risks-group module toggles, the same
    shifting-letters defect W2's slugs exist to stop tests from
    re-introducing.
    """
    out_dir, result = _build_workbook(tmp_path_factory, enabled=enabled)
    tail = (result.stdout + result.stderr)[-4000:]
    assert result.returncode == 0, f"build with {TOGGLE}={enabled} failed:\n{tail}"
    book = out_dir / "retirement_plan.xlsx"
    assert book.exists()

    from openpyxl import load_workbook
    wb = load_workbook(book, read_only=True)
    try:
        present = [n for n in wb.sheetnames if n.endswith("Divorce-QDRO")]
        if enabled:
            assert present, f"expected a Divorce/QDRO tab; sheets were {wb.sheetnames}"
            ws = wb[present[0]]
            text = "\n".join(
                str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None
            )
            assert "DIVORCE / QDRO STRESS TEST" in text
            assert "Headline Comparison" in text
            assert "Asset Split Detail" in text
        else:
            assert not present, f"Divorce/QDRO tab should be absent; sheets were {wb.sheetnames}"
    finally:
        wb.close()

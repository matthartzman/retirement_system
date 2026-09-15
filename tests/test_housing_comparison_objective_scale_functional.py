"""Objective-Value comparability with `10. Social Security` (H12) --
docs/superpowers/plans/2026-09-09-housing-estimate-realism-and-dollar-
convention-design.md, §4.4.

The design's claim is not merely that the housing sweep has *an* objective; it
is that the sweep is "scored like the Social Security sweep" -- same LCV basis
(PV of lifetime consumption plus PV of after-tax terminal transfer, discounted
at ``_roth_discount_rate``), same 0-100 normalization, same feasibility gate.
A housing sheet that quietly scored on, say, nominal terminal net worth would
still render, still rank, and still look right; the only concrete check that
the claim holds is that both sheets, scoring the SAME underlying plan, land on
the same order of magnitude.

Both sheets are built once here against the frozen sample plan, at a
deliberately tiny Monte Carlo path count -- the SS sweep and the housing sweep
are the two most expensive sheets in the workbook, so this file runs each
exactly once and shares the results.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from src.data_io import load_csv, parse_client
from src.plan_config import ensure_engine_config
from src.planning_engines import project
from src.reporting.sheets_strategy import build_sheet10, build_sheet_housing_comparison

from conftest import TEST_INPUT_DIR
from tests.golden_pricing import FROZEN_GOLDEN_MASTER_PRICES, frozen_holdings_prices

ROOT = Path(__file__).resolve().parents[1]

_FAST_MC_SIMS = 8


@pytest.fixture(scope="module")
def both_sheets():
    c = ensure_engine_config(parse_client(load_csv(TEST_INPUT_DIR / "client_data.csv"), ""), source="test")
    # Same config object for both sheets -- the point of the comparison is that
    # they score the same underlying plan, so any difference in magnitude is
    # attributable to the scoring convention, not to the inputs.
    c['roth_policy'] = 'none'
    c['mc_sims'] = _FAST_MC_SIMS
    c['mc_sensitivity_sims'] = 1
    c['housing_sweep_mc_sims'] = _FAST_MC_SIMS
    with frozen_holdings_prices(FROZEN_GOLDEN_MASTER_PRICES):
        rows = project(c)
        ss = build_sheet10(Workbook().create_sheet('ss'), c, rows)
        housing = build_sheet_housing_comparison(Workbook().create_sheet('housing'), c, rows)
    return ss, housing


# @pytest.mark.nightly: engine-internals-only equivalence/sweep-breadth
# check identified in the 2026-09-15 CI-time profiling (see
# documentation/reference/TESTING_REFACTOR_RECOMMENDATIONS.md); cannot be
# triggered by an ordinary UI/config change, so it moved off the PR fast
# tier and runs in the nightly full-suite workflow instead.
@pytest.mark.nightly
def test_both_sheets_actually_produced_scored_candidates(both_sheets):
    ss, housing = both_sheets
    assert ss['scenarios'], "the SS sweep scored no claim-age pairs"
    assert housing is not None and housing['refine_candidates'], \
        "the housing sweep scored no trajectories"


def test_objective_values_are_on_the_same_scale(both_sheets):
    ss, housing = both_sheets
    ss_best = ss['best']['objective_value']
    housing_best = housing['recommended']['objective_value']
    assert ss_best > 0 and housing_best > 0
    ratio = housing_best / ss_best
    # Both are PV-of-consumption-plus-PV-of-after-tax-terminal-transfer on the
    # same plan, so they must agree to well within an order of magnitude. They
    # are NOT expected to be equal: SS's objective adds its survivor-period SS
    # income term (§4.4 drops it -- nothing housing-specific plays that role),
    # and each sweep's winner is a different what-if of the same base plan.
    assert 0.5 < ratio < 2.0, (
        f"housing Objective Value ({housing_best:,.0f}) and Social Security Objective Value "
        f"({ss_best:,.0f}) are not on a comparable scale (ratio {ratio:.2f}) -- the housing "
        "sweep is not scoring on the same LCV basis"
    )


def test_housing_objective_is_the_lcv_score_without_a_survivor_term(both_sheets):
    _ss, housing = both_sheets
    # §4.4: no survivor-income term, so objective_value IS lcv_score here --
    # the one deliberate divergence from build_sheet10's objective.
    for cand in housing['refine_candidates']:
        assert cand['objective_value'] == cand['lcv_score']


def test_both_sheets_normalize_scores_to_the_same_0_100_convention(both_sheets):
    ss, housing = both_sheets
    for candidates in (ss['scenarios'], housing['refine_candidates']):
        scores = [x['rank_score'] for x in candidates]
        assert max(scores) == 100
        assert min(scores) >= 0
        assert all(isinstance(s, int) for s in scores)


def test_both_sheets_gate_on_the_same_feasibility_threshold(both_sheets):
    from src.planning_engines import LCV_FEASIBILITY_GATE_THRESHOLD
    ss, housing = both_sheets
    for candidates in (ss['scenarios'], housing['refine_candidates']):
        for cand in candidates:
            assert cand['feasibility_gate_met'] == (
                cand['feasibility_probability'] >= LCV_FEASIBILITY_GATE_THRESHOLD)

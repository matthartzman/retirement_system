"""Rental lines on the report sheets (WP12.7): Executive Summary, Cash Flow, Net Worth, Balance Sheet,
Lifetime Taxes, Charts and the Results Explorer pages, all only while ``rental_property_income`` is on.

Includes the cash-flow reconciliation (the Cash Bridge Gap): the displayed income columns, rental
cash included, add up to the engine's income and the bridge equals the engine's surplus.
"""
from __future__ import annotations

import json

import pytest
from openpyxl import Workbook, load_workbook

from src.reporting import rental_report_data as rrd
from src.reporting import workbook_common as wc
from src.reporting.sheets_projection_cashflow import build_sheet6
from src.reporting.sheets_projection_charts import build_sheet8
from src.reporting.sheets_projection_net_worth import build_sheet5
from src.reporting.sheets_projection_tax import build_sheet7
from src.reporting.sheets_summary_builder import build_sheet1
from src.reporting.sheets_tax_reporter import build_sheet3
from tests.golden_pricing import frozen_holdings_prices
from tests.rental_sheet_helpers import ALL_CASES, ROOT, built_workbooks, case, header_col, row_labeled, texts

def test_off_adds_no_rental_sheet_columns_or_lines():
    c, rows = case("off")
    assert not rrd.rental_active(rows) and rrd.build_report(c, rows) is None
    ws6 = Workbook().active
    build_sheet6(ws6, c, rows)
    assert "Rental Cash (net)" not in texts(ws6) and "Rental Sale Tax" not in texts(ws6)
    ws5 = Workbook().active
    build_sheet5(ws5, c, rows)
    assert "Rental Equity" not in texts(ws5)
    ws7 = Workbook().active
    build_sheet7(ws7, c, rows)
    assert "Rental Schedule E in AGI" not in texts(ws7)
    with frozen_holdings_prices():
        ws3 = Workbook().active
        build_sheet3(ws3, c, rows)
        ws1 = Workbook().active
        build_sheet1(ws1, c, rows, {}, ss_sweep=None)
    assert not any("Rental" in t for t in texts(ws3) + texts(ws1))
    wb8 = Workbook()
    ws8 = wb8.active
    ws8.title = "1E. Charts"
    build_sheet8(ws8, c, rows)
    data = wb8["_Chart Dashboard Data"]
    assert [data.cell(row=4, column=col).value for col in (27, 29)] == ["Σ Income", "Year"]
    assert not any("Rental" in str(data.cell(row=4, column=col).value) for col in range(1, data.max_column + 1))


# ── cash-flow reconciliation (the Cash Bridge Gap) ──────────────────────────────
@pytest.mark.parametrize("name", ALL_CASES)
def test_cash_flow_sheet_foots_with_rental_cash(name):
    c, rows = case(name)
    ws = Workbook().active
    build_sheet6(ws, c, rows)
    c_rent = header_col(ws, 2, "Rental Cash (net)")
    c_rsale = header_col(ws, 2, "Rental Sale Tax")
    c_inc = header_col(ws, 2, "Σ Income")
    c_gap = header_col(ws, 2, "Reinvested Surplus (forced income > need)")
    c_tcn = header_col(ws, 2, "Total Cash Need")
    c_tax = header_col(ws, 2, "Total Taxes")
    c_draw = header_col(ws, 2, "Elective Portfolio Withdrawals (excl. RMDs)")
    c_inf = header_col(ws, 2, "Income Funding")
    assert header_col(ws, 2, "RMD Dist") + 1 == c_rent and c_rent + 1 == c_inc  # the column sits inside the income group
    income_cols = range(header_col(ws, 2, "Earned"), c_inc)
    for i, r in enumerate(rows):
        line = 3 + i
        bd = r["cashflow_breakdown"]
        val = lambda col: ws.cell(row=line, column=col).value  # noqa: E731
        assert val(c_rent) == pytest.approx(r["rental_net_cash_flow"])
        # the visible income columns (rental cash included) add up to Σ Income, which is the engine's income
        assert sum(val(col) for col in income_cols) == pytest.approx(val(c_inc), abs=0.01)
        assert val(c_inc) == pytest.approx(sum(bd["income"].values()), abs=0.01)
        assert val(c_rsale) == pytest.approx(bd["tax"]["rental_sale"], abs=0.01)
        # the Cash Bridge Gap is the engine's surplus: no rental cash left out of the bridge
        assert -val(c_gap) == pytest.approx(bd["surplus"] - bd["unfunded_gap"], abs=1.0), r["year"]
        assert val(c_inf) + val(c_draw) + val(c_gap) + bd["draws"]["heloc_draw"] == pytest.approx(val(c_tcn), abs=1.0)
        assert val(c_tax) == pytest.approx(r["total_tax"], abs=1.0)


def test_cash_flow_fallback_without_breakdown_counts_rental_cash():
    c, rows = case("profitable")
    stripped = [{k: v for k, v in r.items() if k != "cashflow_breakdown"} for r in rows]
    ws_full, ws_fallback = Workbook().active, Workbook().active
    build_sheet6(ws_full, c, rows)
    build_sheet6(ws_fallback, c, stripped)
    col = header_col(ws_full, 2, "Σ Income")
    for line in range(3, 3 + len(rows)):
        assert ws_fallback.cell(row=line, column=col).value == pytest.approx(ws_full.cell(row=line, column=col).value, abs=1.0)


# ── the other report sheets ─────────────────────────────────────────────────────
@pytest.mark.parametrize("name", ["standalone_loss", "converted_passes"])
def test_net_worth_sheet_shows_rental_equity_and_still_foots(name):
    c, rows = case(name)
    ws = Workbook().active
    build_sheet5(ws, c, rows)
    c_eq = header_col(ws, 2, "Rental Equity")
    c_other = header_col(ws, 2, "Σ Other")
    c_total = header_col(ws, 2, "TOTAL NW")
    assert c_other == c_eq + 1 and c_total == c_other + 1
    for i, r in enumerate(rows):
        line = 3 + i
        assert ws.cell(row=line, column=c_eq).value == pytest.approx(r["rental_equity"])
        assert ws.cell(row=line, column=c_other).value == pytest.approx(r["other_nw"], abs=1.0)
        if r["year"] != c["plan_start"]:  # the plan-start row swaps in opening balances
            assert ws.cell(row=line, column=c_total).value == pytest.approx(r["total_nw"], abs=1.0)
    assert ws.cell(row=1, column=21).value == "OTHER ASSETS"


def test_balance_sheet_lists_rental_value_and_mortgage_and_reconciles():
    c, rows = case("standalone_loss")
    wc.QC_CHECKS.clear()
    with frozen_holdings_prices():
        ws = Workbook().active
        build_sheet3(ws, c, rows)
    labels = texts(ws)
    assert "  Rental Properties" in labels and "  Rental Property Mortgages" in labels
    y0 = rows[0]
    assert ws.cell(row=row_labeled(ws, "Rental Properties"), column=2).value == pytest.approx(y0["rental_property_value"])
    assert ws.cell(row=row_labeled(ws, "Rental Property Mortgages"), column=2).value == pytest.approx(y0["rental_mortgage_balance"])
    check = [q for q in wc.QC_CHECKS if q[0] == "3. Balance Sheet" and "reconciles to projection Y0" in q[1]]
    assert check and check[-1][2] == "PASS", check


def test_lifetime_tax_sheet_adds_rental_columns_and_totals():
    c, rows = case("sale_with_recapture")
    ws = Workbook().active
    build_sheet7(ws, c, rows)
    col_agi = header_col(ws, 4, "Rental Schedule E in AGI")
    col_sale = header_col(ws, 4, "Rental Sale Tax")
    assert col_sale == col_agi + 1 == 18
    assert sum(ws.cell(row=5 + i, column=col_agi).value for i in range(len(rows))) == pytest.approx(
        sum(r["rental_schedule_e_to_agi"] for r in rows))
    sale_year = next(i for i, r in enumerate(rows) if r["year"] == 2035)
    assert ws.cell(row=5 + sale_year, column=col_sale).value == pytest.approx(
        rows[sale_year]["rental_sale_ltcg_tax"] + rows[sale_year]["rental_tax_adjustment"])
    assert any(t.startswith("Lifetime Rental Schedule E added to AGI") for t in texts(ws))
    assert any(t.startswith("Lifetime Tax on Rental Sales") for t in texts(ws))


def test_executive_summary_has_a_rental_block_only_when_on():
    c, rows = case("profitable")
    with frozen_holdings_prices():
        ws = Workbook().active
        build_sheet1(ws, c, rows, {}, ss_sweep=None)
    body = texts(ws)
    assert "Rental Properties" in body
    line = row_labeled(ws, "Rental net cash flow, lifetime (after expenses and debt service)")
    assert ws.cell(row=line, column=2).value == pytest.approx(sum(r["rental_net_cash_flow"] for r in rows))
    assert any("40. Rental Properties" in t for t in body)  # the cross-reference the build rewrites to the final tab name


def _chart_title(chart) -> str:
    t = chart.title
    if t is None:
        return ""
    if isinstance(t, str):
        return t
    return "".join(run.t for para in t.tx.rich.p for run in (para.r or []))


def test_charts_have_rental_series_and_the_income_expense_bars_still_reconcile():
    c, rows = case("standalone_loss")
    wb = Workbook()
    ws = wb.active
    ws.title = "1E. Charts"
    build_sheet8(ws, c, rows)
    data = wb["_Chart Dashboard Data"]
    heads: dict[str, int] = {}
    for col in range(1, data.max_column + 1):
        heads.setdefault(data.cell(row=4, column=col).value, col)
    for h in ("Rental Cash", "Rental Net Cost", "Rental Sale Tax", "Rental equity (right axis)"):
        assert h in heads, h
    c_inc, c_exp = heads["Σ Income"], heads["Σ Spend+Tax"]
    assert heads["Rental Cash"] == c_inc - 1 and heads["Rental Sale Tax"] == c_exp - 1
    for i, r in enumerate(rows):
        line = 5 + i
        bd = r["cashflow_breakdown"]
        net = bd["income"]["rental"]
        assert data.cell(row=line, column=heads["Rental Cash"]).value == max(0, round(net))
        assert data.cell(row=line, column=heads["Rental Net Cost"]).value == max(0, -round(net))
        assert data.cell(row=line, column=heads["Rental Sale Tax"]).value == round(bd["tax"]["rental_sale"])
        if bd["surplus"] < 0 or bd["unfunded_gap"] > 0:
            continue
        # Engine identity: income + draws = spend + tax + surplus. The chart itemizes every tax but the portfolio-LTCG,
        # harvest-credit and 'other' slices, so its two bars differ by exactly that slice (rental cash and sale tax included).
        unitemized = bd["tax"]["ltcg"] + bd["tax"]["tlh_credit"] + bd["tax"]["other"]
        inc, exp = data.cell(row=line, column=c_inc).value, data.cell(row=line, column=c_exp).value
        assert inc - exp == pytest.approx(unitemized, abs=40), r["year"]
    titles = [_chart_title(ch) for ch in ws._charts]
    assert any("Rental Properties" in t for t in titles), titles
    assert any("Rental net cash flow is a series in the Income chart" in t for t in texts(ws))


def test_html_dashboard_reads_the_shifted_chart_columns_when_rental_is_on(tmp_path):
    import json
    import re
    from src.reporting.dashboard import build_html_dashboard
    c, rows = case("standalone_loss")
    wb = Workbook()
    ws = wb.active
    ws.title = "8. Charts Dashboard"
    with frozen_holdings_prices():
        build_sheet8(ws, c, rows)
    xlsx, html_path = tmp_path / "wb.xlsx", tmp_path / "dashboard.out"  # (not a frontend asset)
    wb.save(xlsx)
    with frozen_holdings_prices():
        build_html_dashboard(str(xlsx), str(html_path), rows, c)
    html = html_path.read_text(encoding="utf-8")
    inc_l = json.loads(re.search(r"const INC_L=(\[[^\]]*\]);", html).group(1))
    exp_l = json.loads(re.search(r"const EXP_L=(\[[^\]]*\]);", html).group(1))
    exp_c = json.loads(re.search(r"const EXP_C=(\[[^\]]*\]);", html).group(1))
    inc_c = json.loads(re.search(r"const INC_C=(\[[^\]]*\]);", html).group(1))
    assert inc_l[-1] == "Rental Cash" and exp_l[-2:] == ["Rental Net Cost", "Rental Sale Tax"]
    assert len(exp_c) >= len(exp_l) and len(inc_c) >= len(inc_l)
    years = json.loads(re.search(r"const YEARS=(\[[^\]]*\]);", html).group(1))
    exp_d = json.loads(re.search(r"const EXP_D=(\[.*?\]\]);", html).group(1))
    sale_idx = years.index(2040)
    sale_tax = exp_d[exp_l.index("Rental Sale Tax")][sale_idx]
    assert sale_tax == round(next(r for r in rows if r["year"] == 2040)["cashflow_breakdown"]["tax"]["rental_sale"])


# ── Results Explorer pages (semantic model) ────────────────────────────────────
def _page_table(page: dict, section: int = 0) -> tuple[list[str], list[list]]:
    rows = page["sections"][section]["rows"]
    head = [x["value"] for x in rows[1]["cells"]]
    return head, [[x["value"] for x in r["cells"]] for r in rows[2:]]


@pytest.mark.parametrize("name", ["standalone_loss", "converted_passes"])
def test_explorer_pages_carry_rental_lines_and_foot(name):
    from src.results_model import _balance_sheet_page, _cashflow_page, _chart_page, _net_worth_page
    c, rows = case(name)
    head, data = _page_table(_cashflow_page(c, rows))
    i_rent, i_inc = head.index("Rental Cash (net)"), head.index("Σ Income")
    assert i_inc == i_rent + 1 and head.index("RMD Dist") == i_rent - 1
    for r, line in zip(rows, data):
        assert line[i_rent] == pytest.approx(r["rental_net_cash_flow"])
        assert sum(line[head.index("Earned"):i_inc]) == pytest.approx(line[i_inc], abs=0.01)
    head, data = _page_table(_net_worth_page(c, rows))
    i_eq = head.index("Rental Equity")
    assert head[i_eq + 1:] == ["Σ Other", "TOTAL NW"]
    for r, line in zip(rows, data):
        assert line[i_eq] == pytest.approx(r["rental_equity"])
        assert line[i_eq + 1] + line[i_eq] + line[i_eq - 1] == pytest.approx(r["other_nw"], abs=0.01)  # Σ Other + rental + note
    page = _balance_sheet_page(c, rows)
    values = {x["cells"][0]["value"]: x["cells"][1]["value"] for sec in page["sections"] for x in sec["rows"]}
    assert values["Rental Properties"] == pytest.approx(rows[0]["rental_property_value"])
    assert values["Rental Property Mortgages"] == pytest.approx(rows[0]["rental_mortgage_balance"])
    assert values["NET WORTH"] == pytest.approx(values["Total Assets"] - values["Total Liabilities"])
    chart_page, _ = _chart_page(c, rows, None)
    inc = next(ch for ch in chart_page["charts"] if "Portfolio Draws" in ch["title"])
    exp = next(ch for ch in chart_page["charts"] if "Spending & Taxes" in ch["title"])
    # (the page folds the smallest series into an aggregate past ten, so the rental series may sit inside it;
    # what must hold is that the two bars still reconcile with the rental amounts counted)
    for i in range(min(len(inc["x"]), len(exp["x"]))):  # same tolerance as the no-rental test
        assert abs(sum(x["values"][i] for x in inc["series"]) - sum(x["values"][i] for x in exp["series"])) <= 10


def test_explorer_pages_unchanged_when_off():
    from src.results_model import _balance_sheet_page, _cashflow_page, _chart_page, _net_worth_page
    c, rows = case("off")
    for page in (_cashflow_page(c, rows), _net_worth_page(c, rows), _balance_sheet_page(c, rows)):
        assert "Rental" not in str(page)
    chart_page, _ = _chart_page(c, rows, None)
    assert "Rental" not in str(chart_page)


@pytest.mark.slow
def test_real_workbook_report_sheets_carry_the_rental_lines(tmp_path_factory):
    _off, on_path = built_workbooks(tmp_path_factory)
    on = load_workbook(on_path)
    expected = json.loads((ROOT / "tests" / "fixtures" / "rental_workbook_expectations.json").read_text(encoding="utf-8"))
    for suffix, phrases in expected["other_sheet_phrases"].items():
        other = "\n".join(texts(on[next(n for n in on.sheetnames if n.endswith(suffix))]))
        for phrase in phrases:
            assert phrase in other, f"{suffix} missing {phrase!r}"
    assert "Rental Cash" in "\n".join(texts(on["_Chart Dashboard Data"]))

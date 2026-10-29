"""projection_tables.json (Recent Projections in Financial Trends): the build's Net Worth and Cash Flow
tables as JSON, read back from the worksheets build_sheet5/6 just wrote so the two can't drift."""
from __future__ import annotations

import json

import pytest
from openpyxl import Workbook

from src.reporting.projection_tables import projection_tables_doc, table_from_sheet, write_projection_tables
from src.reporting.sheets_projection_cashflow import build_sheet6
from src.reporting.sheets_projection_net_worth import build_sheet5
from tests.rental_sheet_helpers import ALL_CASES, case


def _sheets(name: str):
    c, rows = case(name)
    ws5, ws6 = Workbook().active, Workbook().active
    return c, {"nw": (ws5, build_sheet5(ws5, c, rows)), "cf": (ws6, build_sheet6(ws6, c, rows))}


@pytest.fixture(scope="module")
def off_sheets():
    return _sheets("off")


@pytest.fixture(scope="module")
def rental_on_sheets():
    return _sheets(next(n for n in ALL_CASES if n != "off"))


@pytest.mark.parametrize("which", ["nw", "cf"])
def test_values_match_sheet_cells(off_sheets, which):
    ws, cols = off_sheets[1][which]
    t = table_from_sheet(ws, cols)
    assert t["rows"]
    for i, row in enumerate(t["rows"]):
        assert row["year"] == ws.cell(3 + i, 1).value
        for key, v in row["values"].items():
            assert v == ws.cell(3 + i, cols[key]).value, (which, row["year"], key)
            assert v is None or (isinstance(v, (int, float)) and not isinstance(v, bool)), (which, key, v)


@pytest.mark.parametrize("which", ["nw", "cf"])
def test_labels_and_banners_are_the_sheet_headers(off_sheets, which):
    ws, cols = off_sheets[1][which]
    by_key = {c["key"]: c for c in table_from_sheet(ws, cols)["columns"]}
    for key, col in cols.items():
        if key in by_key:
            assert by_key[key]["label"] == ws.cell(2, col).value
            assert by_key[key]["banner"]
    first = "Σ Ann" if which == "nw" else "Σ_Inc"
    assert by_key[first]["subtotal"] is True


def test_keys_carry_no_owner_names(off_sheets):
    c, sheets = off_sheets
    names = {str(c.get(k)) for k in ("h_nick", "h_name", "w_nick", "w_name") if c.get(k)}
    assert names
    for ws, cols in sheets.values():
        for col in table_from_sheet(ws, cols)["columns"]:
            assert not any(n in col["key"] for n in names), col


@pytest.mark.parametrize("which", ["nw", "cf"])
def test_totals_rows_excluded_and_ages_split(off_sheets, which):
    ws, cols = off_sheets[1][which]
    t = table_from_sheet(ws, cols)
    assert all(isinstance(r["year"], int) for r in t["rows"])
    assert all(len(r["ages"]) == 2 for r in t["rows"])
    keys = {c["key"] for c in t["columns"]}
    assert not keys & {"Year", "H Age", "W Age", "H_Age", "W_Age"}
    assert t["rows"][0]["ages"] == [ws.cell(3, 2).value, ws.cell(3, 3).value]


def test_conditional_columns_emitted(rental_on_sheets):
    sheets = rental_on_sheets[1]
    nw_keys = {c["key"] for c in table_from_sheet(*sheets["nw"])["columns"]}
    cf_keys = {c["key"] for c in table_from_sheet(*sheets["cf"])["columns"]}
    assert "Rental Equity" in nw_keys
    assert {"Rental", "Rental_Tax"} <= cf_keys


def test_doc_shape(off_sheets):
    sheets = off_sheets[1]
    doc = projection_tables_doc(table_from_sheet(*sheets["nw"]), table_from_sheet(*sheets["cf"]),
                                build_id="b1", built_at="2026-10-03T21:56:40Z", owners=["A", "B"])
    assert doc["schema"] == "projection_tables/1"
    assert doc["build_id"] == "b1" and doc["built_at"] == "2026-10-03T21:56:40Z" and doc["owners"] == ["A", "B"]
    assert set(doc["tables"]) == {"net_worth", "cash_flow"}


def test_write_is_atomic_and_round_trips(tmp_path):
    doc = {"schema": "projection_tables/1", "tables": {"net_worth": {"columns": [], "rows": []}}}
    path = write_projection_tables(tmp_path, doc)
    assert path == tmp_path / "projection_tables.json"
    assert json.loads(path.read_text(encoding="utf-8")) == doc
    assert [p.name for p in tmp_path.iterdir()] == ["projection_tables.json"]


@pytest.mark.slow
def test_real_build_writes_projection_tables_beside_workbook(tmp_path_factory):
    from openpyxl import load_workbook

    from tests.rental_sheet_helpers import built_workbooks

    _off, on_path = built_workbooks(tmp_path_factory)
    doc = json.loads((on_path.parent / "projection_tables.json").read_text(encoding="utf-8"))
    assert doc["schema"] == "projection_tables/1" and doc["built_at"].endswith("Z")
    wb = load_workbook(on_path)
    ws = wb[next(n for n in wb.sheetnames if n.endswith("Net Worth"))]
    total_col = next(c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value == "TOTAL NW")
    first = doc["tables"]["net_worth"]["rows"][0]
    assert first["year"] == ws.cell(3, 1).value
    assert first["values"]["TOTAL NW"] == pytest.approx(ws.cell(3, total_col).value)
    assert "Rental Equity" in first["values"]


def test_unshown_lines_join_the_table_so_groups_foot():
    # The NW sheet's TOTAL subtracts additional liabilities and adds a credit shelter trust, but has
    # no column for either; Recent Projections needs them for its groups to add up to Total NW.
    from src.reporting.projection_tables import add_unshown_nw_lines

    table = {"columns": [{"key": "TOTAL NW", "label": "TOTAL NW", "banner": "TOTAL", "subtotal": True}],
             "rows": [{"year": 2026, "ages": [64, 65], "values": {"TOTAL NW": 10.0}},
                      {"year": 2027, "ages": [65, 66], "values": {"TOTAL NW": 11.0}}]}
    rows = [{"year": 2026, "additional_liabilities": 14286.78, "cst_nw": 0.0},
            {"year": 2027, "additional_liabilities": 0.0, "cst_nw": 0.0}]
    out = add_unshown_nw_lines(table, rows)
    keys = [c["key"] for c in out["columns"]]
    assert "Other_Liab" in keys and "CST" not in keys          # an all-zero line stays out
    liab = next(c for c in out["columns"] if c["key"] == "Other_Liab")
    assert liab["label"] == "Loans & Other Liabilities" and liab["derived"] is True
    assert out["rows"][0]["values"]["Other_Liab"] == -14286.78
    assert out["rows"][1]["values"]["Other_Liab"] == 0.0
    out2 = add_unshown_nw_lines(table, [{"year": 2026, "cst_nw": 5.0}, {"year": 2027, "cst_nw": 6.0}])
    assert [r["values"]["CST"] for r in out2["rows"]] == [5.0, 6.0]
    assert "Other_Liab" not in out2["rows"][0]["values"]


def test_real_plan_groups_foot_after_unshown_lines(off_sheets):
    from src.reporting.projection_tables import add_unshown_nw_lines

    c, sheets = off_sheets
    _c, rows = case("off")
    t = add_unshown_nw_lines(table_from_sheet(*sheets["nw"]), rows)
    for r in t["rows"]:
        v = r["values"]
        shown = sum(v.get(k) or 0 for k in ("Σ Ann", "Σ PreTax", "Σ Roth", "Σ Trust", "HSA_1", "Σ Other", "Other_Liab", "CST"))
        assert abs(shown - v["TOTAL NW"]) < 0.01, (r["year"], shown - v["TOTAL NW"])
